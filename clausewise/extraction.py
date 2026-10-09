"""Document extraction and clause-aware chunking."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Optional

from .config import get_settings
from .errors import DocumentTooLarge, ExtractionFailed, UnsupportedDocument

_CLAUSE_START = re.compile(
    r"^\s*(?:"
    r"\(?(?:\d{1,2}\.\d{1,2}(?:\.\d{1,2})?)\)?"
    r"|\(?[A-Z]\.\)?"
    r"|\(?[ivx]{1,4}\)"
    r"|(?:ARTICLE|Article|CLAUSE|Clause|SECTION|Section)\s+[\dIVXLC]+"
    r")\s+"
)

_DOC_TYPE_HINTS: list[tuple[str, re.Pattern]] = [
    ("murabahah", re.compile(r"\bmuraba(?:hah|ha)\b", re.I)),
    ("ijarah", re.compile(r"\bi?jarah?\b|\bi?jara\b", re.I)),
    ("musharakah", re.compile(r"\bmusharakah?\b|\bshirkah\b", re.I)),
    ("mudarabah", re.compile(r"\bmudarabah?\b", re.I)),
    ("sukuk", re.compile(r"\bsukuk\b", re.I)),
    ("takaful", re.compile(r"\btakaful\b|\btabarru\b", re.I)),
    ("istisna", re.compile(r"\bistisna", re.I)),
    ("salam", re.compile(r"\bsalam\b", re.I)),
    ("qard", re.compile(r"\bqard\b", re.I)),
    ("wakalah", re.compile(r"\bwakalah?\b|\bwakeel\b", re.I)),
    ("trade_facility", re.compile(r"trade\s+facilit|working\s+capital\s+facilit", re.I)),
    ("investment_agreement", re.compile(r"investment\s+agreement|subscription\s+agreement", re.I)),
]

@dataclass
class ExtractedDocument:
    text: str
    page_count: int = 0
    char_count: int = 0
    document_type: str = "unknown"
    warnings: list[str] = field(default_factory=list)
    truncated: bool = False

def _normalise(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)

    lines = text.split("\n")
    rebuilt: list[str] = []
    for line in lines:
        stripped = line.rstrip()
        if not stripped.strip():
            rebuilt.append("")
            continue
        prev = rebuilt[-1] if rebuilt else ""
        if _CLAUSE_START.match(stripped) or re.match(r"^\s*(?:ARTICLE|SECTION)\b", stripped, re.I):
            rebuilt.append(stripped.lstrip())
        elif prev and not prev.endswith((".", ":", ";", "?", "!", "”", '"')) and not _CLAUSE_START.match(prev):
            rebuilt[-1] = f"{prev} {stripped.strip()}"
        else:
            rebuilt.append(stripped.lstrip())

    cleaned = "\n".join(rebuilt)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    cleaned = re.sub(r"[ \t]{2,}", " ", cleaned)
    return cleaned.strip()

_SPECIFIC_INSTRUMENTS = (
    "murabahah",
    "ijarah",
    "musharakah",
    "mudarabah",
    "sukuk",
    "takaful",
    "istisna",
    "salam",
    "qard",
    "wakalah",
)

def detect_document_type(text: str) -> str:
    sample = text[:60_000]
    scores: dict[str, int] = {}
    for name, pattern in _DOC_TYPE_HINTS:
        hits = len(pattern.findall(sample))
        if hits:
            scores[name] = hits
    if not scores:
        return "general"

    specific = {k: v for k, v in scores.items() if k in _SPECIFIC_INSTRUMENTS}
    pool = specific or scores
    return max(pool.items(), key=lambda kv: kv[1])[0]

def extract_from_pdf(content: bytes) -> tuple[str, int, list[str]]:
    warnings: list[str] = []
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ExtractionFailed("That PDF is password-protected. Remove the password and re-upload.")
        pages = reader.pages
        page_count = len(pages)
        limit = get_settings().max_pages_per_document
        chunks: list[str] = []
        empty_pages = 0
        for index, page in enumerate(pages[:limit]):
            try:
                page_text = page.extract_text() or ""
            except Exception:
                page_text = ""
            if not page_text.strip():
                empty_pages += 1
            chunks.append(page_text)
        if page_count > limit:
            warnings.append(
                f"Only the first {limit} pages of {page_count} were analysed. Split the document or raise MAX_PAGES_PER_DOCUMENT."
            )
        if empty_pages and empty_pages >= max(len(chunks) // 2, 1):
            warnings.append(
                "Most pages contained no extractable text — this looks like a scanned document. "
                "Run OCR first; a text-only reading would produce unreliable findings."
            )
        return "\n".join(chunks), page_count, warnings
    except ExtractionFailed:
        raise
    except Exception as exc:
        raise ExtractionFailed(f"We could not read that PDF: {exc}") from exc

def extract_from_docx(content: bytes) -> tuple[str, int, list[str]]:
    warnings: list[str] = []
    try:
        import docx

        document = docx.Document(io.BytesIO(content))
        parts: list[str] = [p.text for p in document.paragraphs if p.text and p.text.strip()]

        table_rows = 0
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
                    table_rows += 1
        if table_rows:
            warnings.append(f"Included {table_rows} table row(s) from the document.")
        if not parts:
            warnings.append("The document contained no readable paragraphs.")
        return "\n".join(parts), 0, warnings
    except Exception as exc:
        raise ExtractionFailed(f"We could not read that Word document: {exc}") from exc

def extract_text_only(content: bytes) -> tuple[str, int, list[str]]:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return content.decode(encoding), 0, []
        except UnicodeDecodeError:
            continue
    raise ExtractionFailed("We could not decode that text file. Save it as UTF-8 and try again.")

def extract_document(content: bytes, *, filename: str = "", mime_type: str = "") -> ExtractedDocument:
    if mime_type == "application/pdf":
        raw, pages, warnings = extract_from_pdf(content)
    elif mime_type.endswith("wordprocessingml.document"):
        raw, pages, warnings = extract_from_docx(content)
    elif mime_type == "text/plain":
        raw, pages, warnings = extract_text_only(content)
    else:
        raise UnsupportedDocument(f"Unsupported document type: {mime_type or 'unknown'}")

    settings = get_settings()
    text = _normalise(raw)

    if not text.strip():
        raise ExtractionFailed(
            "No readable text was found in that document. If it is a scan, run OCR first and upload again."
        )

    truncated = False
    if len(text) > settings.max_chars_per_document:
        text = text[: settings.max_chars_per_document]
        truncated = True
        warnings.append(
            f"Document exceeded {settings.max_chars_per_document:,} characters and was truncated. "
            "Review the tail manually."
        )
    if len(text) < 200:
        warnings.append("The document is very short; findings on a fragment should be treated as indicative only.")

    return ExtractedDocument(
        text=text,
        page_count=pages,
        char_count=len(text),
        document_type=detect_document_type(text),
        warnings=warnings,
        truncated=truncated,
    )

@dataclass
class Chunk:
    index: int
    text: str
    clause_refs: list[str] = field(default_factory=list)
    start: int = 0
    end: int = 0

def chunk_document(
    text: str,
    *,
    target_chars: Optional[int] = None,
    overlap_chars: Optional[int] = None,
) -> list[Chunk]:
    settings = get_settings()
    target = target_chars or settings.chunk_target_chars
    overlap = overlap_chars if overlap_chars is not None else settings.chunk_overlap_chars

    if len(text) <= target:
        refs = [m.group(0).strip() for m in _CLAUSE_START.finditer(text)][:40]
        return [Chunk(index=0, text=text, clause_refs=refs, start=0, end=len(text))]

    boundaries: list[tuple[int, str]] = []
    for match in _CLAUSE_START.finditer(text):
        boundaries.append((match.start(), match.group(0).strip()))

    if len(boundaries) < 3:
        boundaries = [(m.start(), "") for m in re.finditer(r"\n\s*\n", text)]

    if len(boundaries) < 3:
        chunks: list[Chunk] = []
        step = max(target - overlap, 1)
        for i, start in enumerate(range(0, len(text), step)):
            piece = text[start : start + target]
            if not piece.strip():
                continue
            chunks.append(Chunk(index=i, text=piece, start=start, end=start + len(piece)))
        return chunks

    chunks = []
    current_start = 0
    current_refs: list[str] = []
    index = 0

    def flush(end: int) -> None:
        nonlocal index, current_start, current_refs
        piece = text[current_start:end]
        if piece.strip():
            chunks.append(
                Chunk(index=index, text=piece, clause_refs=current_refs[:40], start=current_start, end=end)
            )
            index += 1
        current_refs = []

    for position, ref in boundaries:
        if position - current_start >= target:
            flush(position)
            current_start = max(position - overlap, 0)
        if ref:
            current_refs.append(ref)

    flush(len(text))

    final: list[Chunk] = []
    for chunk in chunks:
        if len(chunk.text) <= target * 2:
            final.append(chunk)
            continue
        sentences = re.split(r"(?<=[.;])\s+", chunk.text)
        buffer = ""
        sub_index = 0
        for sentence in sentences:
            if len(buffer) + len(sentence) > target and buffer:
                final.append(
                    Chunk(
                        index=len(final),
                        text=buffer.strip(),
                        clause_refs=chunk.clause_refs,
                        start=chunk.start,
                        end=chunk.end,
                    )
                )
                buffer = sentence
                sub_index += 1
            else:
                buffer = f"{buffer} {sentence}".strip()
        if buffer.strip():
            final.append(
                Chunk(
                    index=len(final),
                    text=buffer.strip(),
                    clause_refs=chunk.clause_refs,
                    start=chunk.start,
                    end=chunk.end,
                )
            )
    for i, chunk in enumerate(final):
        chunk.index = i
    return final
