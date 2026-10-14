"""Native Word tracked-changes (redline) export.

Put this file at ``clausewise/redline.py``. It needs ``python-docx>=1.2.0`` (for comments).

For every finding that has suggested replacement wording, the module finds the flagged
clause in the document and writes a real Word revision:

* the original wording is wrapped in ``<w:del>`` (shows as struck-through in Word),
* the suggested wording follows in ``<w:ins>`` (shows as underlined),
* a Word comment carrying the rationale and the cited standard is anchored on the change,
* ``<w:trackRevisions/>`` is switched on, so any further edits the reviewer makes are tracked too.

The reviewer accepts or rejects each change in Word. Nothing here decides compliance.

Input is either the original ``.docx`` bytes (edited in place, formatting preserved) or plain
text (a new .docx is built from it). Everything happens in memory: nothing is written to disk
and no document text is stored.

Findings that cannot be located in the document text are not dropped silently. They are
returned in ``RedlineResult.unplaced`` and also appended to the end of the document as a
tracked insertion, so the reviewer still sees them.
"""

from __future__ import annotations

import copy
import io
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any, Iterable, Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_OBJECT = "\ufffc"  # stands in for anything we must not match across (fields, existing revisions)

# Elements that sit between runs and must stop a match from spanning them.
_BARRIER_TAGS = {
    qn("w:ins"),
    qn("w:del"),
    qn("w:moveFrom"),
    qn("w:moveTo"),
    qn("w:hyperlink"),
    qn("w:fldSimple"),
    qn("w:sdt"),
    qn("w:smartTag"),
}
# Run children that mean "whitespace-like, not splittable".
_SPACE_TAGS = {qn("w:tab"), qn("w:br"), qn("w:cr"), qn("w:noBreakHyphen"), qn("w:ptab")}

_CHAR_MAP = {
    0x2018: "'",
    0x2019: "'",
    0x201C: '"',
    0x201D: '"',
    0x00A0: " ",
    0x2013: "-",
    0x2014: "-",
    0x2011: "-",
}

# Children of <w:settings> that must come AFTER <w:trackRevisions> in the schema.
_SETTINGS_SUCCESSORS = {
    qn(f"w:{name}")
    for name in (
        "doNotTrackMoves", "doNotTrackFormatting", "documentProtection", "autoFormatOverride",
        "styleLockTheme", "styleLockQFSet", "defaultTabStop", "autoHyphenation",
        "consecutiveHyphenLimit", "hyphenationZone", "doNotHyphenateCaps", "showEnvelope",
        "summaryLength", "clickAndTypeStyle", "defaultTableStyle", "evenAndOddHeaders",
        "bookFoldRevPrinting", "bookFoldPrinting", "bookFoldPrintingSheets",
        "drawingGridHorizontalSpacing", "drawingGridVerticalSpacing",
        "displayHorizontalDrawingGridEvery", "displayVerticalDrawingGridEvery",
        "doNotUseMarginsForDrawingGridOrigin", "drawingGridHorizontalOrigin",
        "drawingGridVerticalOrigin", "doNotShadeFormData", "noPunctuationKerning",
        "characterSpacingControl", "printTwoOnOne", "strictFirstAndLastChars",
        "noLineBreaksAfter", "noLineBreaksBefore", "savePreviewPicture",
        "doNotValidateAgainstSchema", "saveInvalidXml", "ignoreMixedContent",
        "alwaysShowPlaceholderText", "doNotDemarcateInvalidXml", "saveXmlDataOnly",
        "useXSLTWhenSaving", "saveThroughXslt", "showXMLTags", "alwaysMergeEmptyNamespace",
        "updateFields", "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars",
        "rsids", "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
        "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
        "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
        "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator",
    )
}


# ===================================================================== public types
@dataclass
class RedlineItem:
    """One change to propose: replace ``clause_excerpt`` with ``replacement``."""

    clause_excerpt: str
    replacement: str
    rationale: str = ""
    label: str = ""  # e.g. "Critical · Clause 3.1 · SS-8"


@dataclass
class RedlineResult:
    data: bytes
    applied: list[RedlineItem] = field(default_factory=list)
    unplaced: list[RedlineItem] = field(default_factory=list)


def items_from_findings(findings: Iterable[Any]) -> list[RedlineItem]:
    """Build redline items from `clausewise.models.Finding` rows.

    Referred findings (no cited standard) and findings without replacement wording are skipped:
    there is nothing to propose for them.
    """
    items: list[RedlineItem] = []
    for f in findings:
        severity = (getattr(f, "severity", "") or "").lower()
        if severity == "referred":
            continue
        excerpt = (getattr(f, "clause_excerpt", "") or "").strip()
        wording = (getattr(f, "remedial_wording", "") or "").strip()
        if not excerpt or not wording:
            continue
        label_parts = [
            severity.title() if severity else "",
            f"Clause {f.clause_ref}" if getattr(f, "clause_ref", "") else "",
            getattr(f, "standard_id", "") or "",
        ]
        items.append(
            RedlineItem(
                clause_excerpt=excerpt,
                replacement=wording,
                rationale=(getattr(f, "rationale", "") or "").strip(),
                label=" · ".join(p for p in label_parts if p),
            )
        )
    return items


# ======================================================================= text helpers
def _normalise(text: str) -> tuple[str, list[int]]:
    """Collapse whitespace and unify quotes/dashes. Returns (normalised, map-to-raw-index)."""
    out: list[str] = []
    index: list[int] = []
    previous_space = False
    for i, ch in enumerate(text):
        ch = _CHAR_MAP.get(ord(ch), ch)
        if ch.isspace():
            if previous_space:
                continue
            out.append(" ")
            previous_space = True
        else:
            out.append(ch)
            previous_space = False
        index.append(i)
    return "".join(out), index


# ================================================================= paragraph segments
@dataclass
class _Seg:
    element: Any
    kind: str  # "text" | "space" | "object"
    text: str
    start: int
    end: int


def _segments(paragraph: Any) -> list[_Seg]:
    """Flatten a paragraph into searchable segments. Only direct-child runs are editable."""
    segs: list[_Seg] = []
    cursor = 0
    for child in paragraph:
        tag = child.tag
        if tag == qn("w:r"):
            body = [c for c in child if c.tag != qn("w:rPr")]
            if body and all(c.tag == qn("w:t") for c in body):
                if len(body) > 1:  # coalesce so a run has a single w:t
                    merged = "".join(c.text or "" for c in body)
                    for extra in body[1:]:
                        child.remove(extra)
                    body[0].text = merged
                    body[0].set(_XML_SPACE, "preserve")
                text = body[0].text or ""
                kind = "text"
            elif body and all(c.tag in _SPACE_TAGS for c in body):
                text, kind = " ", "space"
            elif not body:
                continue
            else:
                text, kind = _OBJECT, "object"
        elif tag in _BARRIER_TAGS:
            text, kind = _OBJECT, "object"
            segs.append(_Seg(child, kind, text, cursor, cursor + len(text)))
            cursor += len(text)
            continue
        else:
            continue  # bookmarks, proofing marks, etc. are zero-width
        segs.append(_Seg(child, kind, text, cursor, cursor + len(text)))
        cursor += len(text)
    return segs


def _raw_text(segs: list[_Seg]) -> str:
    return "".join(s.text for s in segs)


def _split_run(seg: _Seg, offset: int) -> None:
    """Split a text run at ``offset`` (relative to the paragraph) into two sibling runs."""
    run = seg.element
    k = offset - seg.start
    if seg.kind != "text" or not (0 < k < len(seg.text)):
        return
    t = run.find(qn("w:t"))
    right = copy.deepcopy(run)
    t.text = seg.text[:k]
    t.set(_XML_SPACE, "preserve")
    right_t = right.find(qn("w:t"))
    right_t.text = seg.text[k:]
    right_t.set(_XML_SPACE, "preserve")
    run.addnext(right)


# =============================================================== revision primitives
class _Ctx:
    def __init__(self, doc: Any, author: str, comments: bool) -> None:
        self.doc = doc
        self.author = author
        self.initials = "".join(w[0] for w in author.split()[:2]).upper() or "CW"
        self.comments = comments
        self.date = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        self._next_id = 9000  # clear of ids already used by the source document
        used = [
            int(v)
            for el in doc.element.body.iter(qn("w:ins"), qn("w:del"))
            if (v := el.get(qn("w:id"), "")).isdigit()
        ]
        if used:
            self._next_id = max(self._next_id, max(used) + 1)

    def rev_id(self) -> str:
        self._next_id += 1
        return str(self._next_id)

    def mark(self, tag: str) -> Any:
        el = OxmlElement(tag)
        el.set(qn("w:id"), self.rev_id())
        el.set(qn("w:author"), self.author)
        el.set(qn("w:date"), self.date)
        return el


def _wrap_deleted(run: Any, ctx: _Ctx) -> Any:
    for t in run.findall(qn("w:t")):
        dt = OxmlElement("w:delText")
        dt.text = t.text
        dt.set(_XML_SPACE, "preserve")
        run.replace(t, dt)
    deletion = ctx.mark("w:del")
    run.addprevious(deletion)
    deletion.append(run)
    return deletion


def _make_inserted_run(text: str, rpr_source: Optional[Any]) -> Any:
    run = OxmlElement("w:r")
    if rpr_source is not None:
        run.append(copy.deepcopy(rpr_source))
    for n, line in enumerate(text.split("\n")):
        if n:
            run.append(OxmlElement("w:br"))
        t = OxmlElement("w:t")
        t.text = line
        t.set(_XML_SPACE, "preserve")
        run.append(t)
    return run


def _add_comment(ctx: _Ctx, first: Any, last: Any, text: str) -> None:
    comment = ctx.doc.comments.add_comment(text=text, author=ctx.author, initials=ctx.initials)
    cid = str(comment.comment_id)
    start = OxmlElement("w:commentRangeStart")
    start.set(qn("w:id"), cid)
    first.addprevious(start)
    end = OxmlElement("w:commentRangeEnd")
    end.set(qn("w:id"), cid)
    last.addnext(end)
    ref_run = OxmlElement("w:r")
    ref = OxmlElement("w:commentReference")
    ref.set(qn("w:id"), cid)
    ref_run.append(ref)
    end.addnext(ref_run)


def _enable_track_revisions(doc: Any) -> None:
    settings = doc.settings.element
    if settings.find(qn("w:trackRevisions")) is not None:
        return
    el = OxmlElement("w:trackRevisions")
    for child in settings:
        if child.tag in _SETTINGS_SUCCESSORS:
            child.addprevious(el)
            return
    settings.append(el)


# ============================================================================ core
def _try_apply(paragraph: Any, item: RedlineItem, ctx: _Ctx) -> bool:
    segs = _segments(paragraph)
    if not segs:
        return False
    raw = _raw_text(segs)
    norm_par, par_index = _normalise(raw)
    norm_ex, _ = _normalise(item.clause_excerpt)
    norm_ex = norm_ex.strip()
    if not norm_ex:
        return False
    pos = norm_par.find(norm_ex)
    if pos < 0:
        return False
    raw_start = par_index[pos]
    raw_end = par_index[pos + len(norm_ex) - 1] + 1

    # Split boundary runs. Do the end first so the start offset stays valid.
    for boundary in (raw_end, raw_start):
        for seg in _segments(paragraph):
            if seg.start < boundary < seg.end:
                _split_run(seg, boundary)
                break

    covered = [
        s for s in _segments(paragraph)
        if s.start >= raw_start and s.end <= raw_end and s.kind in {"text", "space"} and s.end > s.start
    ]
    if not covered:
        return False
    first_run = covered[0].element
    rpr = first_run.find(qn("w:rPr"))

    deletions = [_wrap_deleted(s.element, ctx) for s in covered]
    insertion = ctx.mark("w:ins")
    insertion.append(_make_inserted_run(item.replacement, rpr))
    deletions[-1].addnext(insertion)

    if ctx.comments:
        note = item.label
        if item.rationale:
            note = f"{note}: {item.rationale}" if note else item.rationale
        if note:
            _add_comment(ctx, deletions[0], insertion, note)
    return True


def _append_unplaced(doc: Any, unplaced: list[RedlineItem], ctx: _Ctx) -> None:
    """Add findings we could not locate as a tracked insertion at the end of the document."""

    def tracked_paragraph(text: str, bold: bool = False) -> None:
        para = doc.add_paragraph()
        p = para._p
        ppr = p.get_or_add_pPr()
        rpr_mark = OxmlElement("w:rPr")
        rpr_mark.append(ctx.mark("w:ins"))
        ppr.append(rpr_mark)
        insertion = ctx.mark("w:ins")
        run = _make_inserted_run(text, None)
        if bold:
            rpr = OxmlElement("w:rPr")
            rpr.append(OxmlElement("w:b"))
            run.insert(0, rpr)
        insertion.append(run)
        p.append(insertion)

    tracked_paragraph(
        "Clausewise: findings that could not be matched to the document text. "
        "Review these manually against the source.",
        bold=True,
    )
    for item in unplaced:
        head = f"[{item.label}] " if item.label else ""
        tracked_paragraph(f"{head}Flagged: {item.clause_excerpt}")
        tracked_paragraph(f"Suggested wording: {item.replacement}")
        if item.rationale:
            tracked_paragraph(f"Rationale: {item.rationale}")


def build_redline_docx(
    items: list[RedlineItem],
    *,
    docx_bytes: Optional[bytes] = None,
    text: Optional[str] = None,
    author: str = "Clausewise",
    comments: bool = True,
) -> RedlineResult:
    """Return a .docx with native tracked changes for ``items``.

    Provide ``docx_bytes`` (the uploaded Word file, edited in place) or ``text`` (a new
    document is built from it, one paragraph per line).
    """
    if docx_bytes is None and text is None:
        raise ValueError("Provide docx_bytes or text.")

    if docx_bytes is not None:
        doc = Document(io.BytesIO(docx_bytes))
    else:
        doc = Document()
        for line in (text or "").splitlines():
            doc.add_paragraph(line)

    ctx = _Ctx(doc, author, comments)
    paragraphs = list(doc.element.body.iter(qn("w:p")))

    applied: list[RedlineItem] = []
    unplaced: list[RedlineItem] = []
    for item in items:
        placed = False
        for paragraph in paragraphs:
            if _try_apply(paragraph, item, ctx):
                placed = True
                break
        (applied if placed else unplaced).append(item)

    if unplaced:
        _append_unplaced(doc, unplaced, ctx)
    _enable_track_revisions(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    return RedlineResult(data=buffer.getvalue(), applied=applied, unplaced=unplaced)
